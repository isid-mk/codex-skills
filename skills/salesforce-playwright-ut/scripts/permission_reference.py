"""Resolve the permission workbook from the current reference index, read only."""
import argparse
import hashlib
import json
from pathlib import Path
import sys
import openpyxl


def resolve_permission(index_path):
    index_path = Path(index_path).resolve()
    book = openpyxl.load_workbook(index_path, read_only=True, data_only=True)
    try:
        if '権限設計書' not in book.sheetnames:
            raise ValueError('一覧に「権限設計書」シートがありません')
        sheet = book['権限設計書']
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, ())
        if header.count('資料名') != 1 or header.count('格納先のパス') != 1:
            raise ValueError('権限設計書一覧の見出しを一意に特定できません')
        name_col, path_col = header.index('資料名'), header.index('格納先のパス')
        entries = []
        for n, row in enumerate(rows, 2):
            name, location = row[name_col], row[path_col]
            if not name and not location:
                continue
            if not name or not location:
                raise ValueError(f'権限設計書の資料名またはパス未記載: {n}行')
            path = Path(str(location))
            if not path.is_absolute():
                path = index_path.parent / path
            if path.suffix.lower() in ('.xlsx', '.xlsm'):
                if path.name != Path(str(name)).name:
                    raise ValueError('権限設計書の資料名とファイル名が不一致')
            else:
                path = path / str(name)
            entries.append((path.resolve(), n))
        if len(entries) != 1:
            raise ValueError('権限設計書の参照先は一意に指定してください')
        path, row = entries[0]
        if not path.is_file():
            raise ValueError('権限設計書の参照先が存在しません: ' + str(path))
        return {'path': str(path), 'hash': hashlib.sha256(path.read_bytes()).hexdigest(),
                'indexPath': str(index_path), 'indexHash': hashlib.sha256(index_path.read_bytes()).hexdigest(),
                'sheet': sheet.title, 'row': row, 'networkAccess': False}
    finally:
        book.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--index', required=True)
    args = parser.parse_args()
    try:
        print(json.dumps(resolve_permission(args.index), ensure_ascii=False, indent=2))
        return 0
    except (OSError, ValueError, KeyError) as error:
        print(json.dumps({'ready': False, 'error': str(error)}, ensure_ascii=False), file=sys.stderr)
        return 2


if __name__ == '__main__':
    sys.exit(main())
